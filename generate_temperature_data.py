"""
High-Low Temperature Cycling Test Data Generator

Generates 168 hours (7 days) of continuous temperature cycling data with 21 complete cycles.
Each cycle is 8 hours long with the following phases:
  1. 25°C → -55°C:  30 min
  2. -55°C constant: 120 min
  3. -55°C → 125°C:  90 min
  4. 125°C constant: 120 min
  5. 125°C → 25°C:   60 min
  6. 25°C constant:   60 min
Total per cycle: 480 min = 8 h × 21 cycles = 168 h
"""

import numpy as np
import pandas as pd
from datetime import datetime, timedelta

# ── Configuration ──────────────────────────────────────────────────────────────
SAMPLE_INTERVAL_MIN = 5          # minutes between samples
TOTAL_HOURS = 168                # 7 days
NUM_CYCLES = 21                  # 168 h / 8 h per cycle
TOTAL_SAMPLES = TOTAL_HOURS * 60 // SAMPLE_INTERVAL_MIN  # 2016

ROOM_TEMP = 25.0
LOW_TEMP = -55.0
HIGH_TEMP = 125.0

NOISE_SIGMA = 0.8                # °C sensor noise standard deviation

# Phase durations in minutes for one 8-hour cycle
PHASE_DURATIONS = [30, 120, 90, 120, 60, 60]   # sum = 480 min = 8 h
PHASE_NAMES = [
    "Ramp Down",
    "Hold Low",
    "Ramp Up",
    "Hold High",
    "Ramp to Room",
    "Hold Room",
]

# ── Helpers ────────────────────────────────────────────────────────────────────

def cubic_ease(t: float) -> float:
    """Smooth cubic ease-in-out interpolation (t in [0, 1])."""
    if t <= 0.5:
        return 4 * t * t * t
    else:
        t2 = t - 1
        return 1 + 4 * t2 * t2 * t2


def interpolate(start: float, end: float, t: float) -> float:
    """Interpolate between start and end using cubic easing."""
    return start + (end - start) * cubic_ease(t)


def humidity_for_temp(temp: float) -> float:
    """
    Estimate relative humidity from temperature.
    Low temperatures → higher RH (condensation risk),
    high temperatures → lower RH.
    """
    if temp < LOW_TEMP + 5:
        return 90.0
    elif temp > HIGH_TEMP - 5:
        return 10.0
    else:
        # Linear interpolation across the range
        frac = (temp - LOW_TEMP) / (HIGH_TEMP - LOW_TEMP)
        return 90.0 - 80.0 * frac


# ── Build set-point profile (one sample at a time) ─────────────────────────────

rng = np.random.default_rng(seed=42)  # reproducible noise

timestamps = []
elapsed_hours = []
cycle_numbers = []
phase_labels = []
setpoint_temps = []
actual_temps = []
deviations = []
humidities = []

start_time = datetime(2025, 1, 1, 0, 0, 0)

# Phase boundaries (minutes from cycle start)
phase_boundaries = []
t = 0
for d in PHASE_DURATIONS:
    phase_boundaries.append(t)
    t += d
# phase_boundaries = [0, 30, 150, 240, 360, 420]  (cycle ends at 480)
CYCLE_DURATION_MIN = sum(PHASE_DURATIONS)  # 480

prev_actual = ROOM_TEMP  # exponential-smoothing state

for sample_idx in range(TOTAL_SAMPLES):
    elapsed_min = (sample_idx + 1) * SAMPLE_INTERVAL_MIN  # 5, 10, …, 10080

    # Which cycle (0-based) and position within cycle
    cycle_idx = elapsed_min // CYCLE_DURATION_MIN          # 0 … 20 (edge case: 21 at t=10080)
    pos_in_cycle = elapsed_min % CYCLE_DURATION_MIN        # 0 … 479

    # Edge case: t=10080 is the boundary of cycle 21 → clamp to last sample of cycle 20
    if cycle_idx >= NUM_CYCLES:
        cycle_idx = NUM_CYCLES - 1
        pos_in_cycle = CYCLE_DURATION_MIN  # will resolve to end of Hold Room phase

    # Identify current phase
    phase_idx = 0
    for i in range(len(phase_boundaries) - 1, -1, -1):
        if pos_in_cycle >= phase_boundaries[i]:
            phase_idx = i
            break

    phase_start_min = phase_boundaries[phase_idx]
    phase_dur = PHASE_DURATIONS[phase_idx]
    t_frac = (pos_in_cycle - phase_start_min) / phase_dur  # 0.0 … ≤1.0
    t_frac = max(0.0, min(1.0, t_frac))

    # Set-point temperature for each phase
    if phase_idx == 0:   # Ramp Down: 25 → -55
        setpoint = interpolate(ROOM_TEMP, LOW_TEMP, t_frac)
    elif phase_idx == 1: # Hold Low: -55
        setpoint = LOW_TEMP
    elif phase_idx == 2: # Ramp Up: -55 → 125
        setpoint = interpolate(LOW_TEMP, HIGH_TEMP, t_frac)
    elif phase_idx == 3: # Hold High: 125
        setpoint = HIGH_TEMP
    elif phase_idx == 4: # Ramp to Room: 125 → 25
        setpoint = interpolate(HIGH_TEMP, ROOM_TEMP, t_frac)
    else:                # Hold Room: 25
        setpoint = ROOM_TEMP

    # Simulate actual temperature: exponential lag + Gaussian noise
    alpha = 0.85  # smoothing factor (higher = faster response)
    actual = alpha * setpoint + (1 - alpha) * prev_actual
    actual += rng.normal(0, NOISE_SIGMA)
    prev_actual = actual

    deviation = actual - setpoint
    humidity = humidity_for_temp(setpoint) + rng.normal(0, 1.5)
    humidity = float(np.clip(humidity, 5.0, 98.0))

    ts = start_time + timedelta(minutes=elapsed_min)
    timestamps.append(ts.strftime("%Y-%m-%d %H:%M:%S"))
    elapsed_hours.append(round(elapsed_min / 60, 4))
    cycle_numbers.append(cycle_idx + 1)       # 1-based
    phase_labels.append(PHASE_NAMES[phase_idx])
    setpoint_temps.append(round(setpoint, 2))
    actual_temps.append(round(actual, 2))
    deviations.append(round(deviation, 2))
    humidities.append(round(humidity, 1))

# ── Assemble DataFrame ─────────────────────────────────────────────────────────

df = pd.DataFrame({
    "Timestamp":       timestamps,
    "Elapsed_Hours":   elapsed_hours,
    "Cycle_Number":    cycle_numbers,
    "Phase":           phase_labels,
    "Setpoint_Temp_C": setpoint_temps,
    "Actual_Temp_C":   actual_temps,
    "Deviation_C":     deviations,
    "Humidity_RH":     humidities,
})

# ── Validation ─────────────────────────────────────────────────────────────────

assert len(df) == TOTAL_SAMPLES, f"Expected {TOTAL_SAMPLES} rows, got {len(df)}"
assert df["Elapsed_Hours"].iloc[-1] == TOTAL_HOURS, \
    f"Last Elapsed_Hours should be {TOTAL_HOURS}, got {df['Elapsed_Hours'].iloc[-1]}"
assert df["Cycle_Number"].max() == NUM_CYCLES, \
    f"Expected {NUM_CYCLES} cycles, got {df['Cycle_Number'].max()}"

# ── Export to Excel ────────────────────────────────────────────────────────────

OUTPUT_FILE = "temperature_cycling_test_data.xlsx"

with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Temperature Data")
    wb = writer.book
    ws = writer.sheets["Temperature Data"]

    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Header style
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="CCCCCC")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border

    # Phase colour map
    phase_colors = {
        "Ramp Down":    "DDEEFF",
        "Hold Low":     "BDD7EE",
        "Ramp Up":      "FFE0CC",
        "Hold High":    "FFCCBB",
        "Ramp to Room": "D9EAD3",
        "Hold Room":    "EAF1DD",
    }

    for row_idx in range(2, len(df) + 2):
        phase = ws.cell(row=row_idx, column=4).value
        fill_color = phase_colors.get(phase, "FFFFFF")
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = cell_border

    # Column widths
    col_widths = [20, 14, 13, 15, 17, 15, 13, 13]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

print("🔬 Temperature Cycling Test Data Generator")
print(f"✓ Generated {len(df)} data points")
print(f"✓ Cycles: {df['Cycle_Number'].min()} – {df['Cycle_Number'].max()} ({NUM_CYCLES} total)")
print(f"✓ Time span: 0 – {df['Elapsed_Hours'].iloc[-1]} hours")
print(f"✓ Temperature range: {df['Actual_Temp_C'].min():.1f}°C to {df['Actual_Temp_C'].max():.1f}°C")
print(f"✓ Avg deviation: {df['Deviation_C'].abs().mean():.2f}°C")
print(f"✓ Max deviation: {df['Deviation_C'].abs().max():.2f}°C")
print(f"✓ Avg humidity: {df['Humidity_RH'].mean():.1f}%")
print(f"✓ File saved: {OUTPUT_FILE}")
